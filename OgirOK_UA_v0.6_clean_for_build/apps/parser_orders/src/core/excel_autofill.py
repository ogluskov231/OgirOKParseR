"""Excel autofill for the N template.

Goals:
- take parser records (from docx_parser.py) and write them into the correct sheets
- keep template formatting (styles, data validations, conditional formatting) as much as possible
- be resilient to mapping/schema evolution (old vs new field names)

This module intentionally uses only openpyxl.
"""

from __future__ import annotations

import datetime
import json
import os
from copy import copy
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


DEFAULT_KEY_COL = "E"  # in template: PIB


def _norm_space(s: str) -> str:
    return " ".join((s or "").replace("\u00a0", " ").split()).strip()


def _safe_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _as_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    if isinstance(v, str):
        return v.strip().lower() in {"1", "true", "yes", "так", "y"}
    return False


def _parse_ua_date(v: Any) -> Any:
    """Return datetime.date for dd.mm.yyyy strings, otherwise return original."""
    if isinstance(v, datetime.date):
        return v
    if not isinstance(v, str):
        return v
    s = v.strip()
    if not s:
        return ""
    try:
        parts = s.replace("/", ".").split(".")
        if len(parts) == 3:
            d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
            if y < 100:
                y += 2000
            return datetime.date(y, m, d)
    except Exception:
        return v
    return v


def _get_nested(d: Dict[str, Any], path: str) -> Any:
    cur: Any = d
    for part in path.split("."):
        if not isinstance(cur, dict):
            return None
        if part not in cur:
            return None
        cur = cur[part]
    return cur


def _coalesce(*vals: Any) -> Any:
    for v in vals:
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        return v
    return ""


def _apply_learning(value: Any, mapping: Dict[str, str]) -> Any:
    if not isinstance(value, str):
        return value
    low = value.lower()
    for k, v in (mapping or {}).items():
        if k and k.lower() in low:
            return v
    return value


def _load_json(path: str, default: Any) -> Any:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def route_sheet(record: Dict[str, Any], mapping: Optional[Dict[str, Any]] = None) -> Optional[str]:
    """Decide target sheet from record.

    If mapping provides routing_rules, they are used. Otherwise fallback to hardcoded rules.
    """
    if mapping and isinstance(mapping.get("routing_rules"), list):
        et = _safe_str(record.get("event_type") or record.get("event") or "").lower()
        attached = _as_bool(record.get("is_attached"))
        for rule in mapping["routing_rules"]:
            cond = (rule or {}).get("if") or {}
            if cond.get("event_type") and _safe_str(cond.get("event_type")).lower() != et:
                continue
            if "is_attached" in cond and bool(cond.get("is_attached")) != attached:
                continue
            sh = rule.get("sheet")
            if sh:
                return sh

    # Fallback for older configs
    et = _safe_str(record.get("event_type") or "").lower()
    attached = _as_bool(record.get("is_attached"))
    if et == "arrival" and not attached:
        return "Н 56 омпбр"
    if et == "departure" and not attached:
        return "В 56 омпбр"
    if et == "arrival" and attached:
        return "Н придані 56 омпбр"
    if et == "departure" and attached:
        return "В придані 56 омпбр"
    return None


def _detect_header_row(ws, key_col: str = DEFAULT_KEY_COL, max_scan: int = 25) -> int:
    """Try to detect the header row (where key column contains 'П.І.Б.')."""
    for r in range(1, max_scan + 1):
        v = ws[f"{key_col}{r}"].value
        if isinstance(v, str) and "п.і.б" in v.lower():
            return r
    # fallback to template convention
    return 3


def _detect_last_header_col(ws, header_row: int, max_col: int = 60) -> str:
    last = 5
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if v not in (None, ""):
            last = c
    return get_column_letter(last)


def first_empty_row(ws, key_col: str = DEFAULT_KEY_COL, start_row: int = 4, max_row: int = 5000) -> int:
    """Find first empty row by key column."""
    for r in range(start_row, max_row + 1):
        v = ws[f"{key_col}{r}"].value
        if not _safe_str(v):
            return r
    return max_row + 1


def _copy_cell_style(src: Cell, dst: Cell) -> None:
    # openpyxl style objects are immutable; copy() is safe.
    dst._style = copy(src._style)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)
    dst.comment = copy(src.comment) if src.comment else None


def _copy_cell_formula_or_value(src: Cell, dst: Cell) -> None:
    v = src.value
    if isinstance(v, str) and v.startswith("="):
        try:
            dst.value = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
        except Exception:
            dst.value = v
    elif v is not None and v != "":
        # Only copy constants from template row when they exist.
        dst.value = v


def ensure_row_template(
    ws,
    target_row: int,
    template_row: int,
    start_col: str,
    end_col: str,
    copy_formulas: bool = True,
) -> None:
    """Make sure the target row looks like template_row (styles + optionally formulas).

    This is the missing piece that makes "append to the template" behave like an Excel Table.
    """
    if target_row == template_row:
        return

    def col_to_idx(col: str) -> int:
        from openpyxl.utils import column_index_from_string

        return column_index_from_string(col)

    s = col_to_idx(start_col)
    e = col_to_idx(end_col)
    for c in range(s, e + 1):
        src = ws.cell(row=template_row, column=c)
        dst = ws.cell(row=target_row, column=c)
        _copy_cell_style(src, dst)
        if copy_formulas:
            _copy_cell_formula_or_value(src, dst)

    # row height
    if template_row in ws.row_dimensions and ws.row_dimensions[template_row].height:
        ws.row_dimensions[target_row].height = ws.row_dimensions[template_row].height


def _extend_data_validations(ws, template_row: int, target_row: int) -> None:
    """Extend existing data validations so dropdowns keep working for appended rows.

    Many templates define DV ranges like H3:H999. When we append beyond 999,
    Excel stops showing dropdowns unless ranges are expanded.
    """
    if not getattr(ws, "data_validations", None):
        return

    for dv in list(getattr(ws.data_validations, "dataValidation", []) or []):
        try:
            sq = str(dv.sqref or "").strip()
            if not sq:
                continue
            parts = sq.split()
            changed = False
            out_parts = []
            for part in parts:
                min_col, min_row, max_col, max_row = range_boundaries(part)
                # Only extend ranges that include the template row.
                if min_row <= template_row <= max_row and target_row > max_row:
                    out_parts.append(
                        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{target_row}"
                    )
                    changed = True
                else:
                    out_parts.append(part)
            if changed:
                dv.sqref = " ".join(out_parts)
        except Exception:
            continue


def _extend_tables_if_any(ws, new_last_row: int) -> None:
    """If worksheet contains Excel Tables, extend their ref to include new rows."""
    if not getattr(ws, "tables", None):
        return
    for tbl in ws.tables.values():
        try:
            min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
            if new_last_row > max_row:
                tbl.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_last_row}"
        except Exception:
            continue


def _load_dictionaries_from_template(template_path: str) -> Dict[str, Any]:
    wb = load_workbook(template_path, data_only=True)
    if "Категорії" not in wb.sheetnames:
        return {"units_list": [], "arrival_sources_list": [], "departure_reasons_list": [], "abbr_map": {}}
    ws = wb["Категорії"]
    units: List[str] = []
    arrivals: List[str] = []
    reasons: List[str] = []
    abbr_map: Dict[str, str] = {}

    for r in range(3, 400):
        v = ws[f"A{r}"].value
        if v:
            units.append(str(v).strip())
    for r in range(2, 400):
        v = ws[f"B{r}"].value
        if v:
            arrivals.append(str(v).strip())
    for r in range(2, 400):
        v = ws[f"C{r}"].value
        if v:
            reasons.append(str(v).strip())
    for r in range(1, 800):
        a = ws[f"D{r}"].value
        b = ws[f"E{r}"].value
        if a and b and isinstance(a, str) and isinstance(b, str):
            a2, b2 = a.strip(), b.strip()
            if a2 and b2:
                abbr_map[a2.lower()] = b2
    return {
        "units_list": units,
        "arrival_sources_list": arrivals,
        "departure_reasons_list": reasons,
        "abbr_map": abbr_map,
    }


def _resolve_sheet_mapping(mapping: Dict[str, Any], sheet: str) -> Dict[str, Any]:
    sheets = mapping.get("sheet_mappings") or mapping.get("sheets") or {}
    cur = (sheets.get(sheet) or {}).copy()
    parent = cur.get("inherits")
    if parent:
        base = _resolve_sheet_mapping(mapping, parent)
        # merge: base + override (columns merged shallowly)
        merged = dict(base)
        merged.update({k: v for k, v in cur.items() if k != "columns"})
        cols = dict(base.get("columns") or {})
        cols.update(cur.get("columns") or {})
        merged["columns"] = cols
        return merged
    return cur


def _normalize_record(record: Dict[str, Any], learning: Dict[str, Dict[str, str]], abbr_map: Dict[str, str]) -> Dict[str, Any]:
    """Normalize record to the *new* mapping schema.

    Supports old keys from docx_parser.py.
    """
    r = dict(record or {})

    pib = _norm_space(_safe_str(r.get("pib")))
    parts = pib.split()
    if len(parts) >= 3:
        parts[0] = parts[0].upper()
        parts[1] = parts[1].capitalize()
        parts[2] = parts[2].capitalize()
        pib = " ".join(parts[:3])
    r["pib"] = pib

    # unify order ref
    r["order_ref"] = _coalesce(r.get("order_ref"), r.get("order_number"), r.get("order_no"), "")

    # unify notes
    r["notes"] = _coalesce(r.get("notes"), r.get("note"), "")

    # unify date
    r["event_date"] = _coalesce(
        r.get("event_date"),
        r.get("arrival_date"),
        r.get("departure_date"),
        r.get("date"),
        "",
    )

    # unify subdivisions
    r["to_subdivision"] = _coalesce(r.get("to_subdivision"), r.get("to_unit"), "")
    r["from_subdivision"] = _coalesce(r.get("from_subdivision"), r.get("from_unit"), "")

    # unify from_where
    r["from_where"] = _coalesce(r.get("from_where"), r.get("arrival_source"), "")

    # unify dob
    r["dob"] = _coalesce(r.get("dob"), r.get("birth_date"), "")

    # normalize subdivisions via learning + abbreviation map
    for key in ("to_subdivision", "from_subdivision"):
        if _safe_str(r.get(key)):
            v = _apply_learning(_safe_str(r[key]), learning.get("units", {}))
            low = v.lower()
            for full, abbr in (abbr_map or {}).items():
                if full in low:
                    v = abbr
                    break
            r[key] = v

    if _safe_str(r.get("departure_reason")):
        r["departure_reason"] = _apply_learning(_safe_str(r["departure_reason"]), learning.get("reasons", {}))

    if _safe_str(r.get("arrival_source")):
        r["arrival_source"] = _apply_learning(_safe_str(r["arrival_source"]), learning.get("arrival_sources", {}))

    # flags: support both flat and nested
    flags = r.get("flags") if isinstance(r.get("flags"), dict) else {}
    # Old keys (flat) -> nested
    for old, new in (
        ("has_bzvp", "bzvp"),
        ("fit_support", "fit_support_units"),
        ("needs_vlk", "need_vlk"),
        ("unfit_vlk", "unfit_vlk"),
        ("has_release_grounds", "has_discharge_grounds"),
        ("refuse_tasks", "refuses_combat"),
    ):
        if old in r and new not in flags:
            flags[new] = r.get(old)
    r["flags"] = flags

    return r


def _get_value_for_column(rec: Dict[str, Any], spec: Dict[str, Any]) -> Any:
    """Resolve mapping column spec into an Excel-ready value."""
    field = spec.get("field") or spec.get("parser_field")
    if not field:
        return ""

    # nested/dotted path
    val = _get_nested(rec, field) if isinstance(field, str) and "." in field else rec.get(field)
    if (val is None or val == "") and spec.get("fallback"):
        fb = spec.get("fallback")
        val = _get_nested(rec, fb) if isinstance(fb, str) and "." in fb else rec.get(fb)

    # compose block
    if spec.get("compose") and isinstance(spec.get("compose"), list):
        parts = []
        for p in spec["compose"]:
            lbl = _safe_str(p.get("label"))
            fp = p.get("field")
            pv = _get_nested(rec, fp) if isinstance(fp, str) and "." in fp else rec.get(fp)
            if _safe_str(pv):
                parts.append(f"{lbl}: {_safe_str(pv)}" if lbl else _safe_str(pv))
        if parts:
            base = _safe_str(val)
            return (base + "\n" if base else "") + "\n".join(parts)

    typ = spec.get("type")
    if typ == "flag":
        if _as_bool(val) or (_safe_str(val).lower() == "так"):
            return spec.get("true_value") or "Так"
        return ""
    if typ == "date":
        return _parse_ua_date(val)
    return val


def autofill_template(
    records: List[Dict[str, Any]],
    template_path: str,
    output_dir: str,
    rules_path: str,
    mapping_path: str,
    diagnostics_path: Optional[str] = None,
    learning_dir: str = "config/learning",
    max_rows_per_sheet: int = 5000,
) -> str:
    """Fill the Excel template and save to output_dir.

    Notes:
    - rules_path is kept for compatibility (GUI references it), but mapping drives placement.
    - This function does NOT modify the template file on disk.
    """
    os.makedirs(output_dir, exist_ok=True)

    wb = load_workbook(template_path)
    _ = _load_json(rules_path, {})  # kept for compatibility
    mapping = _load_json(mapping_path, {})

    dicts = _load_dictionaries_from_template(template_path)
    units_list = dicts["units_list"]
    arrival_list = dicts["arrival_sources_list"]
    reasons_list = dicts["departure_reasons_list"]
    abbr_map = dicts["abbr_map"]

    learning: Dict[str, Dict[str, str]] = {}
    for name in ("units", "reasons", "arrival_sources"):
        p = os.path.join(learning_dir, f"{name}.json")
        learning[name] = _load_json(p, {}) if os.path.exists(p) else {}

    unknowns: Dict[str, Any] = {
        "unknown_units": [],
        "unknown_departure_reasons": [],
        "unknown_arrival_sources": [],
        "skipped": [],
    }

    # Pre-compute per-sheet header + column span for style copy
    sheet_layout: Dict[str, Tuple[int, int, str, str]] = {}
    for sheet in wb.sheetnames:
        if sheet == "Категорії":
            continue
        ws = wb[sheet]
        header_row = _detect_header_row(ws)
        data_row = header_row + 1
        last_col = _detect_last_header_col(ws, header_row)
        sheet_layout[sheet] = (header_row, data_row, "A", last_col)

    for rec0 in records:
        rec = _normalize_record(rec0, learning, abbr_map)
        sheet = route_sheet(rec, mapping=mapping)
        if not sheet or sheet not in wb.sheetnames:
            unknowns["skipped"].append({"reason": "no_sheet", "rec": rec})
            continue

        ws = wb[sheet]
        shmap = _resolve_sheet_mapping(mapping, sheet)
        key_col = ((shmap.get("key") or {}).get("column") or DEFAULT_KEY_COL).strip().upper()

        header_row, data_start_row, start_col, end_col = sheet_layout.get(sheet, (3, 4, "A", "T"))
        row = first_empty_row(ws, key_col=key_col, start_row=data_start_row, max_row=max_rows_per_sheet)

        # ensure formatting for this row
        ensure_row_template(ws, target_row=row, template_row=data_start_row, start_col=start_col, end_col=end_col)
        _extend_data_validations(ws, template_row=data_start_row, target_row=row)
        _extend_tables_if_any(ws, new_last_row=row)

        col_specs: Dict[str, Dict[str, Any]] = shmap.get("columns") or {}
        for xl_col, spec in col_specs.items():
            xl_col = xl_col.strip().upper()
            if not xl_col or not spec:
                continue
            val = _get_value_for_column(rec, spec)
            if val is None or val == "":
                continue
            ws[f"{xl_col}{row}"].value = val

        # unknown diagnostics based on dictionaries
        et = _safe_str(rec.get("event_type")).lower()
        if et == "arrival":
            if _safe_str(rec.get("arrival_source")) and rec["arrival_source"] not in arrival_list:
                unknowns["unknown_arrival_sources"].append(rec["arrival_source"])
            if _safe_str(rec.get("to_subdivision")) and rec["to_subdivision"] not in units_list:
                unknowns["unknown_units"].append(rec["to_subdivision"])
        elif et == "departure":
            if _safe_str(rec.get("departure_reason")) and rec["departure_reason"] not in reasons_list:
                unknowns["unknown_departure_reasons"].append(rec["departure_reason"])
            if _safe_str(rec.get("from_subdivision")) and rec["from_subdivision"] not in units_list:
                unknowns["unknown_units"].append(rec["from_subdivision"])

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(output_dir, f"N_filled_{ts}.xlsx")
    wb.save(out_path)

    if diagnostics_path:
        os.makedirs(os.path.dirname(diagnostics_path), exist_ok=True)
        for k in ("unknown_units", "unknown_departure_reasons", "unknown_arrival_sources"):
            unknowns[k] = sorted(list({x for x in unknowns[k] if x}))
        with open(diagnostics_path, "w", encoding="utf-8") as f:
            json.dump(unknowns, f, ensure_ascii=False, indent=2)

    return out_path


# -----------------------------
# Mapping generator (etalon)
# -----------------------------

def ensure_mapping_file(mapping_path: str, template_sources_path: str) -> None:
    """
    Ensure config/autofill/mapping.json exists.
    If missing, regenerate it from the official templates defined in template_sources.json.

    This is intentionally conservative: it will NOT overwrite an existing mapping.
    """
    import json as _json
    from pathlib import Path as _Path

    mp = _Path(mapping_path)
    if mp.exists():
        return

    ts = _Path(template_sources_path)
    if not ts.exists():
        raise FileNotFoundError(f"Missing template_sources.json: {template_sources_path}")

    sources = _json.loads(ts.read_text(encoding="utf-8"))
    project_root = ts.parent.parent.parent  # .../config/autofill -> project root

    # prefer local generator script if present
    gen_py = project_root / "scripts" / "generate_mapping_from_etalon.py"
    if gen_py.exists():
        # run as a module-less script to keep dependencies minimal
        import runpy as _runpy
        _runpy.run_path(str(gen_py), run_name="__main__")
        return

    raise RuntimeError("Cannot regenerate mapping: generator script missing.")


